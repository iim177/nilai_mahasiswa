"""
File ini berisi fungsi-fungsi untuk:
- Membuat template Excel kosong (untuk diisi & di-import)
- Membaca file Excel yang diupload (import mahasiswa)
- Membuat file Excel hasil export (transkrip nilai)
"""

import io
from openpyxl import Workbook, load_workbook


def buat_template_mahasiswa() -> io.BytesIO:
    """Membuat file Excel kosong dengan header NPM dan Nama, untuk di-download & diisi user."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Mahasiswa"
    ws.append(["NPM", "Nama"])
    ws.append(["2021010001", "Contoh Nama Mahasiswa"])

    # Lebarkan kolom biar enak dibaca
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 35

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def baca_excel_mahasiswa(file_bytes: bytes):
    """
    Membaca file Excel yang diupload user.
    Return list of dict: [{"npm": "...", "nama": "..."}, ...]
    Baris pertama dianggap header dan dilewati.
    """
    wb = load_workbook(io.BytesIO(file_bytes))
    ws = wb.active

    hasil = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # lewati baris header
        if row is None or len(row) < 2:
            continue
        npm, nama = row[0], row[1]
        if npm is None or nama is None:
            continue
        hasil.append({"npm": str(npm).strip(), "nama": str(nama).strip()})
    return hasil


def buat_excel_transkrip(kelas, anggota_list) -> io.BytesIO:
    """
    Membuat file Excel transkrip nilai untuk satu kelas.
    anggota_list = list of KelasMahasiswa (sudah include relasi mahasiswa)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Transkrip Nilai"

    ws.append([f"Transkrip Nilai - {kelas.mata_kuliah.nama_makul} ({kelas.nama_kelas})"])
    ws.append([f"Tahun Ajaran: {kelas.tahun_ajaran.nama_tahun_ajaran}"])
    ws.append([f"Dosen: {kelas.dosen.nama_lengkap if kelas.dosen else '-'}"])
    ws.append([])
    header = ["No", "NPM", "Nama"]
    if kelas.pakai_bobot:
        header += [
            f"Tugas ({kelas.bobot_tugas:.0f}%)",
            f"Kuis ({kelas.bobot_kuis:.0f}%)",
            f"UTS ({kelas.bobot_uts:.0f}%)",
            f"UAS ({kelas.bobot_uas:.0f}%)",
            "Nilai Akhir",
        ]
    else:
        header += ["Tugas", "Kuis", "UTS", "UAS"]
    ws.append(header)

    for i, am in enumerate(anggota_list, start=1):
        baris = [i, am.mahasiswa.npm, am.mahasiswa.nama, am.nilai_tugas_final(), am.nilai_kuis_final(), am.nilai_uts, am.nilai_uas]
        if kelas.pakai_bobot:
            baris.append(am.nilai_akhir())
        ws.append(baris)

    lebar = [5, 18, 30, 12, 12, 12, 12] + ([12] if kelas.pakai_bobot else [])
    for col, width in zip("ABCDEFGH", lebar):
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
